# -*- coding: utf-8 -*-
"""점검 과정에서 발견한 결함들에 대한 회귀 테스트."""

import io
import os
import sys
import threading
import time
import unittest
from datetime import date

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE, "python-lib"))

from primepatent.config import PATENT_TERM_YEARS, ScoringConfig  # noqa: E402
from primepatent.export import export_bytes, safe_cell  # noqa: E402
from primepatent.family import build_families, representative_records  # noqa: E402
from primepatent.jobs import JobManager  # noqa: E402
from primepatent.llm.analyzer import LLMCache  # noqa: E402
from primepatent.parsing import to_date  # noqa: E402
from primepatent.records import build_records  # noqa: E402
from primepatent.scoring.context import AnalysisContext  # noqa: E402
from primepatent.scoring.engine import score_records  # noqa: E402
from primepatent.status import remaining_term_years  # noqa: E402

AS_OF = date(2026, 8, 25)

MAPPING = {key: {"column": key} for key in
           ["applicationNumber", "country", "legalStatus", "familyId", "familyMembers",
            "claimCount", "earliestPriorityDate", "forwardCitationCount"]}


def make_rows():
    """6개국 패밀리 1건 + 단일국 패밀리 5건.

    비교집단을 모집단 전체로 잡으면 6개국 패밀리의 값이 분포를 지배한다.
    """
    rows = []
    members = ";".join("%s102020000001" % c for c in ["KR", "US", "JP", "CN", "EP", "TW"])
    for country in ["KR", "US", "JP", "CN", "EP", "TW"]:
        rows.append({"applicationNumber": "%s102020000001" % country, "country": country,
                     "legalStatus": "공개", "familyId": "BIG", "familyMembers": members,
                     "claimCount": 30, "earliestPriorityDate": "2020-01-01",
                     "forwardCitationCount": 50})
    for index in range(5):
        rows.append({"applicationNumber": "KR10202000000%d" % (index + 2), "country": "KR",
                     "legalStatus": "공개", "familyId": "S%d" % index,
                     "familyMembers": "KR10202000000%d" % (index + 2),
                     "claimCount": 5 + index, "earliestPriorityDate": "2020-01-01",
                     "forwardCitationCount": index})
    return rows


class PeerPopulationTest(unittest.TestCase):
    """백분위 비교집단은 채점 단위(패밀리 대표문헌)와 같아야 한다."""

    def setUp(self):
        self.config = ScoringConfig(topic_name="T", peer_min_size=2, llm_enabled=False)
        self.records = build_records(make_rows(), MAPPING)
        build_families(self.records, self.config, AS_OF)
        self.targets = representative_records(self.records, {}, True)

    def test_peer_group_uses_scoring_unit(self):
        self.assertEqual(len(self.records), 11)
        self.assertEqual(len(self.targets), 6)      # 패밀리 6개
        ctx = AnalysisContext(self.records, self.config, AS_OF, peer_records=self.targets)
        self.assertEqual(len(ctx.peer_records), 6)
        # 인용 해석용 색인은 전체 문헌을 계속 사용해야 한다
        self.assertGreater(len(ctx.doc_index), 6)

    def test_score_records_does_not_let_one_family_dominate(self):
        scored = score_records(self.targets, {}, self.config, AS_OF, population=self.records)
        for row in scored["rows"]:
            components = {c["key"]: c for a in row["areas"].values() for c in a["components"]}
            # 백분위 비교집단 표본 수 = 대표문헌 수(6). 11 이면 패밀리 중복이 섞인 것.
            peer_n = components["impact.leadership"]["detail"]["peerN"]
            self.assertEqual(peer_n, 6,
                             "%s 의 비교집단 표본이 채점 단위와 다릅니다: %d" % (row["docNumber"], peer_n))
            # 외부TR/TR 최대값을 구하는 모집단도 같은 단위여야 한다.
            for key in ("impact.competitorCoverage", "impact.citation"):
                population_n = components[key]["detail"]["populationN"]
                self.assertEqual(population_n, 6,
                                 "%s 의 %s 모집단이 채점 단위와 다릅니다: %d"
                                 % (row["docNumber"], key, population_n))


class RemainingTermTest(unittest.TestCase):
    """존속기간은 제도상 상한(20년)을 넘을 수 없다."""

    def test_future_priority_date_is_capped(self):
        value = remaining_term_years({"earliestPriorityDate": to_date("2099-01-01")}, AS_OF)
        self.assertEqual(value, PATENT_TERM_YEARS)

    def test_bogus_expiry_date_is_capped(self):
        value = remaining_term_years({"expiryDate": to_date("2999-01-01")}, AS_OF)
        self.assertEqual(value, PATENT_TERM_YEARS)

    def test_normal_case_unchanged(self):
        value = remaining_term_years({"earliestPriorityDate": to_date("2019-01-10")}, AS_OF)
        self.assertAlmostEqual(value, 12.38, places=1)


class JobResultRetentionTest(unittest.TestCase):
    """완료된 작업 결과를 무한 보관하면 백엔드가 메모리 부족으로 죽는다."""

    def test_old_results_are_released(self):
        manager = JobManager(result_retention=2)
        jobs = [manager.submit("t", (lambda n: (lambda job: {"n": n}))(i)) for i in range(5)]
        deadline = time.time() + 5
        while time.time() < deadline and any(job.finished_at is None for job in jobs):
            time.sleep(0.05)
        retained = [job for job in jobs if job.result is not None]
        self.assertEqual(len(retained), 2)
        dropped = [job for job in jobs if job.result_dropped]
        self.assertEqual(len(dropped), 3)
        self.assertTrue(all(job.status == "done" for job in jobs))   # 상태는 유지
        self.assertTrue(dropped[0].to_dict()["resultDropped"])


class LLMCacheTest(unittest.TestCase):
    """공유 캐시는 스스로 락을 갖고 크기 상한이 있어야 한다."""

    def test_bounded_and_thread_safe(self):
        cache = LLMCache(max_entries=100)
        errors = []

        def worker(index):
            try:
                for i in range(400):
                    cache.put("k%d" % (i % 150), {"v": index})
                    cache.get("k%d" % (i % 150))
            except Exception as exc:      # noqa: BLE001
                errors.append(exc)

        threads = [threading.Thread(target=worker, args=(i,)) for i in range(8)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        self.assertEqual(errors, [])
        self.assertLessEqual(len(cache), 100)

    def test_returns_copy_not_reference(self):
        cache = LLMCache()
        cache.put("k", {"score": 1})
        first = cache.get("k")
        first["score"] = 999
        self.assertEqual(cache.get("k")["score"], 1)


class MarketEntryTest(unittest.TestCase):
    """주요 시장 진입도는 '실제 출원이 있는 국가' 만 세야 한다."""

    MAP = {k: {"column": k} for k in
           ["applicationNumber", "country", "legalStatus", "familyMembers",
            "familyCountryDocCounts", "designatedStates"]}

    def _entry(self, row):
        from primepatent.scoring.context import AnalysisContext
        from primepatent.scoring.engine import score_one
        config = ScoringConfig(peer_min_size=2)
        record = build_records([row], self.MAP)[0]
        build_families([record], config, AS_OF)
        ctx = AnalysisContext([record], config, AS_OF, peer_records=[record])
        return [c for area in score_one(record, None, ctx)["areas"].values()
                for c in area["components"] if c["key"] == "market.entry"][0]

    def test_designated_states_do_not_count_as_filings(self):
        """PCT/EPC 지정국은 '출원 가능한 나라' 목록일 뿐 실제 출원이 아니다.

        회귀: 지정국을 포함하면 US 단독 출원 1건이 만점(15점)을 받았다.
        """
        component = self._entry({
            "applicationNumber": "US10475776", "country": "US", "legalStatus": "등록(존속)",
            "familyMembers": "US10475776", "familyCountryDocCounts": "US:1",
            "designatedStates": "AE;AT;AU;CN;DE;EP;JP;KR;GB;FR;TW"})
        self.assertEqual(component["detail"]["countries"], ["US"])
        self.assertAlmostEqual(component["score"], 4.6, places=3)   # 소비 4.2 + 공급망 0.4

    def test_full_score_requires_all_weighted_countries(self):
        component = self._entry({
            "applicationNumber": "KR1", "country": "KR", "legalStatus": "등록(존속)",
            "familyMembers": "KR1;US2;JP3;CN4;EP5;TW6",
            "familyCountryDocCounts": "KR:2|US:3|JP:1|CN:1|EP:1|TW:1", "designatedStates": ""})
        self.assertEqual(component["score"], 15.0)

    def test_missing_ep_cannot_reach_full_score(self):
        """EP 가 없으면 소비시장 2.1점을 받을 수 없으므로 만점이 불가능하다."""
        component = self._entry({
            "applicationNumber": "KR1", "country": "KR", "legalStatus": "등록(존속)",
            "familyMembers": "KR1;US2;JP3;CN4;TW5",
            "familyCountryDocCounts": "KR:2|US:3|JP:1|CN:1|TW:1", "designatedStates": ""})
        self.assertNotIn("EP", component["detail"]["countries"])
        self.assertAlmostEqual(component["score"], 12.9, places=3)
        self.assertLess(component["score"], 15.0)

    def test_country_sources_are_reported(self):
        component = self._entry({
            "applicationNumber": "KR1", "country": "KR", "legalStatus": "등록(존속)",
            "familyMembers": "KR1;US2", "familyCountryDocCounts": "KR:2|US:3",
            "designatedStates": ""})
        sources = component["detail"]["countrySources"]
        self.assertEqual(sources["KR"], "개별국 문헌 수")
        self.assertEqual(sources["US"], "개별국 문헌 수")


class CountryParserTest(unittest.TestCase):
    """국가코드는 독립 토큰일 때만 인식해야 한다."""

    def test_normal_formats(self):
        from primepatent.parsing import country_doc_counts
        self.assertEqual(country_doc_counts("KR:2|US:3|JP:1"), {"KR": 2, "US": 3, "JP": 1})
        self.assertEqual(country_doc_counts("KR(2), US(3)"), {"KR": 2, "US": 3})
        self.assertEqual(country_doc_counts("KR;US;JP"), {"KR": 1, "US": 1, "JP": 1})
        self.assertEqual(country_doc_counts("kr:2|us:3"), {"KR": 2, "US": 3})

    def test_words_do_not_produce_phantom_countries(self):
        """회귀: EPO→EP, Europe→RO, SEPTEMBER→SE/PT/BE 로 오인되어 점수가 부풀려졌다."""
        from primepatent.parsing import country_doc_counts
        self.assertEqual(country_doc_counts("EPO패밀리 기준 KR:2 US:3"), {"KR": 2, "US": 3})
        self.assertNotIn("EP", country_doc_counts("EPO패밀리 기준"))
        self.assertEqual(country_doc_counts("US:3 (Europe 미출원)"), {"US": 3})
        self.assertEqual(country_doc_counts("SEPTEMBER"), {})
        self.assertEqual(country_doc_counts("DEPT KR:1"), {"KR": 1})

    def test_blank(self):
        from primepatent.parsing import country_doc_counts
        self.assertEqual(country_doc_counts(""), {})
        self.assertEqual(country_doc_counts(None), {})


class LlmRationaleExportTest(unittest.TestCase):
    """LLM 판단 근거가 엑셀에 함께 나와야 한다."""

    def _payload(self):
        row = {
            "rank": 1, "applicationNumber": "KR1", "title": "테스트", "docNumber": "KR1",
            "totalScore": 50.0, "grade": "C", "areaScores": {}, "gate": {}, "notes": [],
            "familyCountries": [], "registrationLabel": "등록",
            "areas": {"tech": {"components": [
                {"key": "tech.coreCentrality", "label": "중심성", "score": 4.0, "max": 5.0,
                 "llmScore": 4.0, "detail": {}, "notes": []}]}},
            "llm": {"status": "ok", "provider": "dataiku",
                    "rationale": "종합 근거입니다",
                    "centralityRationale": "핵심 구성이 독립항에 있습니다",
                    "expansionRationale": "관련 청구항 비율 60%",
                    "diversityRationale": "장치와 제조방법 2개 유형",
                    "keyFeatures": ["하이브리드 본딩"],
                    "claimAnalysis": {"coreElementsInIndependentClaims": ["구리 패드"],
                                      "linkageClaimed": True, "relatedClaimCount": 6,
                                      "totalClaimCount": 10, "relatedClaimRatio": 0.6,
                                      "expansionDependentCount": 4,
                                      "claimTypes": ["장치", "제조방법"],
                                      "designAroundRisk": "low"}},
        }
        return {"rows": [row], "summary": {}, "config": {}, "warnings": [], "mapping": {}}

    def test_score_sheet_has_rationale_columns(self):
        from primepatent.export import flatten_row, result_columns
        columns = result_columns()
        for label in ("중심성 근거", "확장도 근거", "유형 다양성 근거", "LLM 종합근거",
                      "독립항 내 핵심 구성요소", "독립항 유형", "관련 청구항 비율"):
            self.assertIn(label, columns)
        flat = flatten_row(self._payload()["rows"][0])
        self.assertEqual(flat["중심성 근거"], "핵심 구성이 독립항에 있습니다")
        self.assertEqual(flat["독립항 유형"], "장치, 제조방법")
        self.assertEqual(flat["연결관계 청구"], "Y")

    def test_dedicated_llm_sheet(self):
        import openpyxl
        from primepatent.export import export_bytes
        book = openpyxl.load_workbook(io.BytesIO(export_bytes(self._payload(), "xlsx")))
        self.assertIn("LLM판단근거", book.sheetnames)
        sheet = book["LLM판단근거"]
        header = [c.value for c in sheet[1]]
        self.assertIn("중심성 근거", header)
        self.assertIn("핵심기술 중심성", header)
        values = [c.value for c in sheet[2]]
        self.assertEqual(values[header.index("핵심기술 중심성")], 4.0)
        self.assertEqual(values[header.index("중심성 근거")], "핵심 구성이 독립항에 있습니다")
        self.assertEqual(values[header.index("관련 청구항 비율")], 0.6)


class ExportInjectionTest(unittest.TestCase):
    """업로드 데이터가 엑셀 수식으로 실행되면 안 된다."""

    def test_safe_cell_prefixes_formula_characters(self):
        for value in ("=1+1", "+CMD", "-1+1", "@SUM(1)"):
            self.assertTrue(safe_cell(value).startswith("'"), value)
        self.assertEqual(safe_cell("정상 제목"), "정상 제목")
        self.assertEqual(safe_cell(None), None)
        self.assertEqual(safe_cell(12), 12)

    def test_xlsx_has_no_formula_cells(self):
        import openpyxl
        row = {"rank": 1, "docNumber": "KR1", "title": '=HYPERLINK("http://evil","x")',
               "applicant": "+CMD", "currentAssignee": "-1+1", "statusLabel": "@SUM(1)",
               "totalScore": 50.0, "grade": "C", "areaScores": {}, "areas": {},
               "gate": {}, "llm": {}, "notes": [], "familyCountries": []}
        payload = {"rows": [row], "summary": {}, "config": {}, "warnings": [], "mapping": {}}
        book = openpyxl.load_workbook(io.BytesIO(export_bytes(payload, "xlsx")))
        formulas = [(sheet.title, cell.coordinate)
                    for sheet in book.worksheets
                    for line in sheet.iter_rows() for cell in line
                    if cell.data_type == "f"]
        self.assertEqual(formulas, [], "수식으로 해석되는 셀이 있습니다: %s" % formulas)


if __name__ == "__main__":
    unittest.main(verbosity=2)
